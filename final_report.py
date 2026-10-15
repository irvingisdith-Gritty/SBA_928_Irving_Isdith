
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

# SBA 928 - Final Project Report
OUTPUT = Path("SBA_928_Final_Report.docx")

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

doc.add_heading("SBA 928 - Enhancing Market Research with AI", 0)
doc.add_paragraph("Fine-Tuning FLAN-T5 for Movie Market Research")
doc.add_paragraph("Student: Irving Isdith")
doc.add_paragraph("October 2026")

def section(title, paragraphs):
    doc.add_heading(title, level=1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)

section("1. Project Overview", [
    "This project explores how fine-tuning a pretrained language model can support movie market research. The selected model was Google's FLAN-T5-small.",
    "The objective was to generate useful explanations of movie popularity, ratings, audience engagement, market positioning, and promotional opportunities."
])

section("2. Dataset and Preparation", [
    "The project used the Top Movies dataset, containing 9,826 complete movie records.",
    "A Python preparation script generated 600 instruction-response examples across six market-research task families. The examples were divided into 480 training examples and 120 validation examples using a random seed of 42.",
    "Each example contained an instruction with movie metadata and a target analytical response. Input and target token lengths were audited; none exceeded the configured 128-token limit."
])

doc.add_heading("3. Market Research Task Families", level=1)

families = {
    "Consumer Behavior": "Interpret popularity, voting activity, and ratings while distinguishing engagement from satisfaction.",
    "Genre Trends": "Analyze genre positioning without claiming a market-wide trend from one movie.",
    "International Markets": "Evaluate language information without assuming geographic audience location.",
    "Audience Segmentation": "Propose content-interest groups without inventing viewer demographics.",
    "Competitor Comparison": "Identify relevant comparison metrics and missing competitive data.",
    "Recommendation Strategy": "Suggest cautious promotional experiments and methods for validating results."
}

for name, description in families.items():
    doc.add_paragraph(f"{name}: {description}", style="List Bullet")

section("4. Model and Training Configuration", [
    "The base model was google/flan-t5-small, loaded using Hugging Face Transformers.",
    "The three-epoch experiment used a batch size of 2, learning rate of 0.00005, CPU training, maximum input length of 128 tokens, and maximum target length of 128 tokens.",
    "The training script used Seq2SeqTrainer, Seq2SeqTrainingArguments, and DataCollatorForSeq2Seq. The model and tokenizer were saved after training."
])

doc.add_heading("5. Training Results", level=1)
doc.add_paragraph(
    "Validation loss decreased during the three-epoch experiment:"
)

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "Epoch"
table.rows[0].cells[1].text = "Validation Loss"

for epoch, loss in [(1, 1.5218), (2, 0.9288), (3, 0.7733)]:
    cells = table.add_row().cells
    cells[0].text = str(epoch)
    cells[1].text = f"{loss:.4f}"

doc.add_paragraph(
    "Validation loss declined by approximately 49.2% between epochs 1 and 3. "
    "This indicates improved prediction of held-out reference tokens, "
    "but does not by itself establish better analytical generation."
)

section("6. Model Comparison and Evaluation", [
    "The project compared the original pretrained FLAN-T5 model with multiple fine-tuned checkpoints. A four-model comparison was saved in model_comparison_results_4models.csv, using 18 independent prompts across six task families.",
    "Additional qualitative tests evaluated the three-epoch checkpoint on analytical prompts, short questions, and a confirmed training example.",
    "The fine-tuned model reproduced several movie statistics correctly and answered some simple questions, but also generated repetitive sentences, incomplete explanations, and unsupported conclusions.",
    "For example, a genre-trend question produced an unsupported claim instead of recognizing that one movie cannot establish a genre-wide trend. A training example about audience engagement reproduced the movie statistics but failed to explain the distinction between engagement and satisfaction.",
    "These results demonstrate that lower validation loss did not consistently translate into stronger independent analytical responses."
])

section("7. Responsible AI, Bias, and Limitations", [
    "Movie popularity and vote counts are limited indicators of attention or engagement and do not directly establish purchases or viewing motivations.",
    "Average ratings provide limited reception information, but can be affected by the number and representativeness of voters.",
    "Original language does not identify where viewers live. Inferring audience nationality, age, income, or other demographic characteristics from movie metadata would be unsupported.",
    "Genre-wide trends require multiple titles and comparable observations over time. Competitive comparisons require consistent measurement definitions.",
    "The evaluation identified factual reliability and repetition problems. These observations highlight the need for human review, additional testing across task families, and careful treatment of uncertain conclusions.",
    "A formal demographic fairness audit was not performed because the available movie metadata does not establish actual viewer demographics."
])

section("8. Reproducibility and Deliverables", [
    "The project includes the source movie dataset, generated fine-tuning dataset, data-preparation script, training scripts, model-comparison script, comparison CSV files, README, and saved model checkpoints.",
    "The three-epoch training experiment used a seeded 80/20 split. Model outputs and validation-loss measurements are documented for reproducibility.",
    "The project was developed using Python, pandas, PyTorch, Hugging Face Transformers, and Hugging Face Datasets."
])

section("9. Reflection and Future Improvements", [
    "The project demonstrated how instruction-response data can be prepared for fine-tuning a pretrained sequence-to-sequence model.",
    "The main challenge was improving analytical response quality rather than merely reducing validation loss.",
    "Future improvements could include more varied reasoning examples, a larger instruction-tuned model, task-specific evaluation rubrics, controlled generation experiments, and additional testing on unseen movies.",
    "Future work should also evaluate subgroup performance using appropriate data and avoid making demographic claims that the dataset cannot support."
])

section("10. Conclusion", [
    "The FLAN-T5 fine-tuning pipeline was implemented successfully and validation loss improved across three epochs.",
    "However, qualitative testing revealed continuing weaknesses in reasoning, factual reliability, and response consistency.",
    "The central lesson is that successful training and lower validation loss are not sufficient measures of real-world model usefulness. Independent generation tests and responsible interpretation remain essential."
])

doc.save(OUTPUT)
print(f"Report created: {OUTPUT.resolve()}")
print(f"File size: {OUTPUT.stat().st_size:,} bytes")
